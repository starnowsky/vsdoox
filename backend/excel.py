"""Формирование таблицы связей в формате Excel (.xlsx)."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = ['Источник', 'Назначение', 'Протокол/порт', 'Описание']
_WIDTHS = [30, 30, 16, 44]


def _safe(value):
    return '' if value is None else str(value)


def build_connections_xlsx(rows):
    """Возвращает байты .xlsx-файла. rows — список строк [источник, назначение, протокол/порт, описание]."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Связи'

    ws.append(HEADERS)
    header_fill = PatternFill('solid', fgColor='DDEBF7')
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for row in rows:
        ws.append([_safe(x) for x in row])

    for idx, width in enumerate(_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=(cell.column == 4))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
